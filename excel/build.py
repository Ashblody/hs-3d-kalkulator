from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList

EUR='#,##0.00 "€"'; PCT='0.00 %'; N2='#,##0.00'
ORANGE=PatternFill('solid',fgColor='FFF1E0'); OB=Side(style='thin',color='F59E0B')
HEAD=PatternFill('solid',fgColor='0F172A'); GREEN='059669'; BLUE='0284C7'
thin=Side(style='thin',color='E2E8F0')
wb=Workbook()

def header(ws,title,widths):
    ws.sheet_view.showGridLines=False
    ws['A1']=title; ws['A1'].font=Font(bold=True,size=16,color='0F172A')
    for k,v in widths.items(): ws.column_dimensions[k].width=v
def inp(c,fmt=None):
    c.fill=ORANGE; c.border=Border(left=OB,right=OB,top=OB,bottom=OB); c.font=Font(color='9A3412',bold=True)
    if fmt: c.number_format=fmt
def thead(ws,row,cols):
    for i,t in enumerate(cols,1):
        c=ws.cell(row,i,t); c.fill=HEAD; c.font=Font(bold=True,color='FFFFFF'); c.alignment=Alignment(wrap_text=True,vertical='center')

# ---------- Tiskalniki ----------
tp=wb.active; tp.title='Tiskalniki'
header(tp,'Tiskalniki',{'A':28,'B':12,'C':12,'D':14,'E':12,'F':9,'G':9,'H':9,'I':9,'J':9,'K':10,'L':14})
types=['PLA','PETG','ASA','PC','TPU','PETG-CF']
thead(tp,3,['Naziv','Nabavna cena','Življ. doba (h)','Servis v dobi','Moč privzeto (W)']+[f'{t}' for t in types]+['Amortizacija €/h'])
tp.cell(2,6,'Moč po materialu (W) – prazno = privzeto').font=Font(italic=True,color='64748B')
printers=[('Bambu Lab X2D Combo',870.4,3000,200,450,{'ASA':450}),('Prusa CORE One INDX',2249,3000,100,None,{}),
          ('Elegoo Mars 2 Pro',450,3000,200,None,{}),('Formlabs Form Studio 1',1650,3000,200,None,{})]
for i,(n,p,l,s,w,pw) in enumerate(printers):
    r=4+i; tp.cell(r,1,n)
    for col,v,f in [(2,p,EUR),(3,l,'#,##0'),(4,s,EUR),(5,w,'0')]:
        c=tp.cell(r,col,v); inp(c,f)
    for j,t in enumerate(types):
        c=tp.cell(r,6+j,pw.get(t)); inp(c,'0')
    tp.cell(r,12,f'=IF(C{r}>0,(B{r}+D{r})/C{r},0)').number_format='0.0000 "€/h"'
for r in range(8,16):
    for col in range(1,12): inp(tp.cell(r,col))
    tp.cell(r,12,f'=IF(A{r}="","",IF(C{r}>0,(B{r}+D{r})/C{r},0))').number_format='0.0000 "€/h"'
tp['A17']='Moči so ocene, dokler jih ne izmeriš. Prazne vrstice 8–15 so za nove tiskalnike.'; tp['A17'].font=Font(italic=True,color='64748B')

# ---------- Materiali ----------
MATEND=45
mt=wb.create_sheet('Materiali')
header(mt,'Materiali',{'A':24,'B':16,'C':10,'D':12,'E':13,'F':13,'G':12,'H':12,'I':14})
thead(mt,3,['Naziv','Proizvajalec','Vrsta','Neto masa (g)','Redna cena koluta','Akcijska / bulk cena','Uporabi','€/kg','Temp. sušenja','Akcija / bulk','Vir cene'])
mats=[('Bambu PLA Basic','Bambu Lab','PLA',13),('Bambu PETG HF','Bambu Lab','PETG',13),('Bambu ASA','Bambu Lab','ASA',26.64),
('Prusament PLA','Prusa','PLA',32),('Prusament PETG','Prusa','PETG',32),('Prusament ASA','Prusa','ASA',32),
('Plastika Trček PLA','Plastika Trček','PLA',21),('Plastika Trček PETG','Plastika Trček','PETG',22),
('Prusament PC Blend','Prusa','PC',46.39),('Plastika Trček TPU','Plastika Trček','TPU',38.57),('Prusament PETG-CF','Prusa','PETG-CF',60)]
import re as _re
_seed=open(__file__.rsplit('/',2)[0]+'/src/seed.ts',encoding='utf8').read()
_base={'BB':'https://eu.store.bambulab.com/products/','PR':'https://www.prusa3d.com/de/kategorie/','TR':'https://3dtrcek.com/sl/'}
mats=[(m[0],m[1],m[2],m[3],1000) for m in mats]
mats2=[]
for g in _re.findall(r"x\('[^']+', '([^']+)', '([^']+)', '([^']+)', (\d+), ([\d.]+), (null|\['([^']+)', ([\d.]+)\]), (BB|PR|TR) \+ '([^']+)'(?:, '([^']+)')?\)",_seed):
    mats2.append(dict(n=g[0],mk=g[1],t=g[2],g=int(g[3]),p=float(g[4]),sl=g[6],sp=float(g[7]) if g[7] else None,src=_base[g[8]]+g[9],dry=g[10] or None))
assert len(mats2)==22, len(mats2)
dvU=DataValidation(type='list',formula1='"Redna,Akcijska"',allow_blank=False); mt.add_data_validation(dvU)
rows=[[m[0],m[1],m[2],m[4],m[3],None,'Redna',None,None,None,None] for m in mats]+[[d['n'],d['mk'],d['t'],d['g'],d['p'],d['sp'],'Redna',None,d['dry'],d['sl'] or None,d['src']] for d in mats2]
mt.column_dimensions['J'].width=14; mt.column_dimensions['K'].width=48
for r in range(4,MATEND+1):
    vals=rows[r-4] if r-4<len(rows) else [None]*11
    for col in (1,2,3,4,5,6,7,9,10,11):
        c=mt.cell(r,col,vals[col-1]); inp(c,EUR if col in (5,6) else ('#,##0' if col==4 else None))
    dvU.add(f'G{r}')
    mt.cell(r,8,f'=IF(OR(A{r}="",D{r}<=0),"",IF(AND(G{r}="Akcijska",ISNUMBER(F{r})),F{r},E{r})/(D{r}/1000))').number_format=EUR

# ---------- Cenik ----------
cn=wb.create_sheet('Cenik storitev')
header(cn,'Cenik storitev',{'A':38,'B':14,'C':8})
rates=[('3D-skeniranje',20,'€/h'),('Urejanje skena',15,'€/h'),('CAD-modeliranje',20,'€/h'),('Zahtevnejši reverse engineering',25,'€/h'),
('Priprava tiska in dodelava',15,'€/h'),('Prevoz',0.30,'€/km'),('Minimalno naročilo',15,'€'),('Privzeti pribitek tiska',100,'%'),
('Privzeta rezerva za neuspele tiske',10,'%'),('Prijateljski popust (samo storitve)',15,'%'),('Prijatelj: cena na uro tiskanja (namesto pribitka)',1,'€/h'),('Predaja datoteke, ko so storitve že v projektu',0,'€'),('Ponovna prodaja datoteke',30,'% prvotne')]
for i,(l,v,u) in enumerate(rates):
    r=3+i; cn.cell(r,1,l); c=cn.cell(r,2,v); inp(c,N2); cn.cell(r,3,u)
R={k:f"'Cenik storitev'!$B${3+i}" for i,k in enumerate(['scan','edit','cad','rev','prep','km','min','markup','reserve','friend','friendH','handover','repeat'])}

# ---------- Poslovne ----------
ps=wb.create_sheet('Poslovne nastavitve')
header(ps,'Poslovne nastavitve',{'A':40,'B':14,'C':10})
bs=[('Ciljna osebna postavka',15,'€/h'),('Obračunljive ure na mesec',20,'h'),('Prispevki in administracija',138.01,'€/mes'),
('Cena elektrike',0.20,'€/kWh'),('Davčna rezerva (od prihodka)',0,'%'),('Obračunavam DDV',"NE",'DA/NE'),('Stopnja DDV',22,'%')]
for i,(l,v,u) in enumerate(bs):
    r=3+i; ps.cell(r,1,l); c=ps.cell(r,2,v); inp(c,None if isinstance(v,str) else N2); ps.cell(r,3,u)
dvDN=DataValidation(type='list',formula1='"DA,NE"'); ps.add_data_validation(dvDN); dvDN.add('B8')
ps['A11']='Davčna rezerva je privzeto 0 %, ker je odvisna od statusa (normiranec …). Vnesi svojo oceno.'; ps['A11'].font=Font(italic=True,color='64748B')
B={k:f"'Poslovne nastavitve'!$B${3+i}" for i,k in enumerate(['target','hours','fixed','kwh','tax','vat','vatpct'])}
# mesečni načrt
ps['A13']='Mesečni načrt po ceniku'; ps['A13'].font=Font(bold=True,size=12)
plan=[('Povprečna urna postavka',f"=AVERAGE({R['scan']},{R['edit']},{R['cad']},{R['rev']},{R['prep']})",EUR),
('Prihodek dela na mesec','=B3*B14'.replace('B3','B4'),EUR),('− prispevki in administracija','=B5',EUR),('− ocena davka','=B15*B7/100',EUR),
('Predvideni osebni zaslužek','=B15-B16-B17',EUR),('Na uro','=IF(B4>0,B18/B4,0)','0.00 "€/h"'),('Opozorilo','=IF(B19<B3,"Cenik ne doseže ciljne postavke!","Cilj dosežen")',None)]
for i,(l,f,fm) in enumerate(plan):
    r=14+i; ps.cell(r,1,l); c=ps.cell(r,2,f)
    if fm: c.number_format=fm
ps['B18'].font=Font(bold=True,color=GREEN); ps['B20'].font=Font(bold=True,color='EA580C')

# ---------- Datoteke ----------
dt=wb.create_sheet('Datoteke')
header(dt,'Prodaja digitalnih datotek',{'A':42,'B':9,'C':9,'D':11,'E':13,'F':13,'G':10})
thead(dt,3,['Datoteka','Vklop','Ure','Postavka €/h','Dodatni stroški','Znesek','Ure (vklop)'])
dig=[('Surov 3D-sken (OBJ/STL)','scan'),('Očiščen in popravljen STL','edit'),('CAD-model izvožen v STL','cad'),
('Parametrični CAD-model, STEP','cad'),('Rekonstrukcija STEP iz skena','rev'),('Celoten paket STL + STEP + dokumentacija','cad')]
dt.add_data_validation(dvDN2:=DataValidation(type='list',formula1='"DA,NE"'))
for i,(l,rk) in enumerate(dig):
    r=4+i; dt.cell(r,1,l)
    c=dt.cell(r,2,'NE'); inp(c); dvDN2.add(f'B{r}')
    inp(dt.cell(r,3,0),N2); inp(dt.cell(r,4,f"={R[rk]}"),EUR); inp(dt.cell(r,5,0),EUR)
    dt.cell(r,6,f'=IF(B{r}="DA",C{r}*D{r}+E{r},0)').number_format=EUR
    dt.cell(r,7,f'=IF(B{r}="DA",C{r},0)').number_format=N2
dt['A10']='Skupaj'; dt['A10'].font=Font(bold=True)
dt['F10']='=SUM(F4:F9)'; dt['F10'].number_format=EUR; dt['F10'].font=Font(bold=True,color=GREEN)
dt['G10']='=SUM(G4:G9)'; dt['G10'].number_format=N2
dt['A12']='Pravice'; dt['A12'].font=Font(bold=True,size=12)
for i,(l,v) in enumerate([('Samo fizični izdelek, brez datotek','DA'),('Datoteka za osebno uporabo naročnika','NE'),('Pravica do komercialne uporabe','NE'),('Prenos izvorne CAD-datoteke','NE')]):
    r=13+i; dt.cell(r,1,l); c=dt.cell(r,2,v); inp(c); dvDN2.add(f'B{r}')
dt['A20']='Samodejna cena datoteke iz ur storitev (Kalkulator C26:C29)'; dt['A20'].font=Font(bold=True,size=12)
dvF=DataValidation(type='list',formula1='"Brez,STL,STEP"'); dt.add_data_validation(dvF)
dvC=DataValidation(type='list',formula1='"Z naročilom,Samostojno,Ponovna"'); dt.add_data_validation(dvC)
for r,l,v in [(21,'Vrsta datoteke (Brez / STL / STEP)','Brez'),(22,'Prodaja (Z naročilom / Samostojno / Ponovna)','Z naročilom'),(23,'Ročna cena € (prazno = predlog)',None)]:
    dt.cell(r,1,l); c=dt.cell(r,2,v); inp(c, EUR if r==23 else None)
dvF.add('B21'); dvC.add('B22')
for r,l,f in [(21,'STL = sken + urejanje',f"=Kalkulator!$C$26*{R['scan']}+Kalkulator!$C$27*{R['edit']}"),
 (22,'STEP = STL + CAD + zahtevni CAD',f"=F21+Kalkulator!$C$28*{R['cad']}+Kalkulator!$C$29*{R['rev']}"),
 (23,'Prvotna cena izbrane datoteke','=IF(B21="STL",F21,IF(B21="STEP",F22,0))'),
 (24,'Predlog',f'=IF(B21="Brez",0,IF(B22="Z naročilom",{R["handover"]},IF(B22="Ponovna",F23*{R["repeat"]}/100,F23)))'),
 (25,'CENA DATOTEKE (pred popustom stopnje)','=IF(AND(B21<>"Brez",ISNUMBER(B23)),B23,F24)'),
 (26,'Sken/urejanje ni obračunano kot storitev (1 = da)','=IF(AND(B21<>"Brez",B22<>"Z naročilom"),1,0)'),
 (27,'CAD ni obračunan kot storitev (1 = da)','=IF(AND(F26=1,B21="STEP"),1,0)')]:
    dt.cell(r,5,l); c=dt.cell(r,6,f); c.number_format=EUR if r<26 else '0'
dt['F25'].font=Font(bold=True,color=GREEN)
dt['A28']='Z naročilom: storitve so že obračunane, datoteka zastonj ali samo predaja. Samostojno/Ponovna: ure skena/CAD so v ceni datoteke in se ne obračunajo še kot storitve.'; dt['A28'].font=Font(italic=True,color='64748B')
dt['A18']='Postavka STEP iz CAD (vrstica 7) in rekonstrukcija iz skena (vrstica 8) sta ločeni. Pravice vplivajo na ceno le prek ur in dodatnih stroškov.'; dt['A18'].font=Font(italic=True,color='64748B')

# ---------- Kalkulator ----------
k=wb.create_sheet('Kalkulator',0)
header(k,'HS 3D Kalkulator',{'A':2,'B':34,'C':24,'D':3,'E':34,'F':16,'G':16,'H':16})
k['A1']=None; k['B1']='HS 3D Kalkulator'; k['B1'].font=Font(bold=True,size=20,color='0F172A')
k['B2']='Oranžna polja spreminjaš, ostalo se izračuna samo.'; k['B2'].font=Font(italic=True,color='64748B')
inputs=[(4,'Ime projekta','Nosilci lučk – 2 leva + 2 desna',None),(5,'Stranka','Demo',None),(6,'Datum','2026-10-25',None),
(7,'Način obračuna','Stranka',None),(8,'Vrsta naročila','Kombinirano',None),
(10,'Tiskalnik','Bambu Lab X2D Combo',None),(11,'Material','Bambu ASA',None),(12,'Število kosov',4,'0'),(13,'Teža kosa (g)',192,N2),
(14,'Čas kosa – ure',4,N2),(15,'Čas kosa – minute',42,N2),(16,'Vnesem skupen čas serije','NE',None),(17,'Skupen čas – ure',0,N2),(18,'Skupen čas – minute',0,N2),
(19,'Dodatna poraba (podpore, purge, AMS) g',0,N2),(20,'Potrošni material (lepilo …) €',3.4,EUR),(21,'Ostali neposredni stroški tiska €',0,EUR),
(22,'Rezerva za neuspele tiske %',f"={R['reserve']}",N2),(23,'Pribitek na stroške tiska %',f"={R['markup']}",N2),(24,'Ročni stroški tiska € (prazno = izračun)',None,EUR),
(26,'3D-skeniranje na terenu (h)',1,N2),(27,'Urejanje skena (h)',1,N2),(28,'CAD-modeliranje (h)',0.75,N2),(29,'Reverse engineering (h)',0,N2),
(30,'Priprava tiska (h)',0.25,N2),(31,'Dodelava / podpore (h)',0.5,N2),(32,'Dodatno delo (h)',0,N2),(33,'Dodatno delo – cena €/h',f"={R['prep']}",EUR),
(34,'Prijateljski popust na storitve %',f"={R['friend']}",N2),
(36,'Prevoz (km)',0,N2),(37,'Pakiranje €',0,EUR),(38,'Dostava €',0,EUR),(39,'Ostali stroški €',0,EUR),(40,'Popust €',0,EUR)]
for sec,row in [('Podatki o projektu',3),('3D-tiskanje',9),('Storitve',25),('Ostalo',35)]:
    k.cell(row,2,sec).font=Font(bold=True,size=12,color=BLUE)
for r,l,v,f in inputs:
    k.cell(r,2,l); c=k.cell(r,3,v); inp(c,f)
from datetime import date
k['C6']=date(2026,10,25); k['C6'].number_format='DD.MM.YYYY'
def dv(formula,cells):
    d=DataValidation(type='list',formula1=formula,allow_blank=False); k.add_data_validation(d); d.add(cells)
dv('"Stranka,Prijatelj,Zase"','C7'); dv('"Samo tisk,Skeniranje,CAD,Kombinirano,Digitalna datoteka"','C8')
dv('=Tiskalniki!$A$4:$A$15','C10'); dv('=Materiali!$A$4:$A$45','C11'); dv('"DA,NE"','C16')
nn=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0',showErrorMessage=True,error='Vnesi število 0 ali več.'); k.add_data_validation(nn)
for r in [13,14,15,17,18,19,20,21,22,23,24,26,27,28,29,30,31,32,33,34,36,37,38,39,40]: nn.add(f'C{r}')
pc=DataValidation(type='whole',operator='greaterThanOrEqual',formula1='1',showErrorMessage=True,error='Vsaj 1 kos.'); k.add_data_validation(pc); pc.add('C12')

Z='$C$7="Zase"'
calc=[
(12,'Čas tiska skupaj (h)','=IF($C$16="DA",$C$17+$C$18/60,$C$12*($C$14+$C$15/60))',N2),
(13,'Filament skupaj (kg)','=($C$12*$C$13+$C$19)/1000','0.000'),
(14,'Vrsta materiala','=INDEX(Materiali!$C:$C,MATCH($C$11,Materiali!$A:$A,0))',None),
(15,'Cena materiala €/kg','=INDEX(Materiali!$H:$H,MATCH($C$11,Materiali!$A:$A,0))',EUR),
(16,'Vrstica tiskalnika','=MATCH($C$10,Tiskalniki!$A:$A,0)','0'),
(17,'Moč (W)','=IFERROR(IF(ISNUMBER(INDEX(Tiskalniki!$F:$K,F16,MATCH(F14,Tiskalniki!$F$3:$K$3,0))),INDEX(Tiskalniki!$F:$K,F16,MATCH(F14,Tiskalniki!$F$3:$K$3,0)),N(INDEX(Tiskalniki!$E:$E,F16))),N(INDEX(Tiskalniki!$E:$E,F16)))','0'),
(18,'Amortizacija €/h','=INDEX(Tiskalniki!$L:$L,F16)','0.0000'),
(20,'Material','=F13*F15',EUR),
(21,'Elektrika',f"=F12*F17/1000*{B['kwh']}",EUR),
(22,'Amortizacija tiskalnika','=F12*F18',EUR),
(23,'Potrošni + ostali neposredni','=$C$20+$C$21',EUR),
(24,'Stroški pred rezervo','=SUM(F20:F23)',EUR),
(25,'Rezerva','=F24*$C$22/100',EUR),
(26,'STROŠKI TISKANJA','=IF(ISNUMBER($C$24),$C$24,F24+F25)',EUR),
(27,'Pribitek % (samo Stranka)','=IF($C$7="Stranka",$C$23,0)',N2),
(28,'Pribitek € (Prijatelj: ure tiska × €/h)',f"=F26*F27/100+IF($C$7=\"Prijatelj\",F12*{R['friendH']},0)",EUR),
(29,'Bruto marža tiska','=IF(F30>0,F28/F30,0)',PCT),
(30,'Prodajna cena tiska','=F26+F28',EUR),
(32,'Storitve po ceniku',f"=($C$26*{R['scan']}+$C$27*{R['edit']})*(1-Datoteke!$F$26)+($C$28*{R['cad']}+$C$29*{R['rev']})*(1-Datoteke!$F$27)+$C$30*{R['prep']}+$C$31*{R['prep']}+$C$32*$C$33",EUR),
(33,'Popust prijatelj %','=IF($C$7="Prijatelj",$C$34/100,0)',PCT),
(34,'Popust / ni obračunano (Zase)',f'=IF({Z},F32,F32*F33)',EUR),
(35,'Storitve (obračunano)','=F32-F34',EUR),
(36,'Digitalne datoteke',f'=IF({Z},0,(Datoteke!$F$10+Datoteke!$F$25)*(1-F33))',EUR),
(37,'Prevoz',f"=$C$36*{R['km']}",EUR),
(38,'Pakiranje + dostava + ostalo','=$C$37+$C$38+$C$39',EUR),
(39,'Popust',f'=IF({Z},0,$C$40)',EUR),
(40,'Vmesna vsota','=F30+F35+F36+F37+F38-F39',EUR),
(41,'Minimalno naročilo upoštevano',f"=AND(NOT({Z}),F40>0,F40<{R['min']})",None),
(42,'PRIHODEK brez DDV',f"=MAX(0,IF(F41,{R['min']},F40))",EUR),
(43,'DDV',f"=IF(AND({B['vat']}=\"DA\",NOT({Z})),F42*{B['vatpct']}/100,0)",EUR),
(44,'KONČNA CENA (z DDV, če je obračunan)','=F42+F43',EUR),
(45,'Cena na kos','=IF($C$12>0,F44/$C$12,F44)',EUR),
(47,'Ure dela','=SUM($C$26:$C$32)+Datoteke!$G$10-IF(Datoteke!$B$22="Ponovna",Datoteke!$F$26*($C$26+$C$27)+Datoteke!$F$27*($C$28+$C$29),0)',N2),
(48,'Vrednost mojega časa (informativno)','=F32+Datoteke!$F$10+Datoteke!$F$25',EUR),
(49,'Dejanski stroški','=F26+F37+F38',EUR),
(50,'Delež fiksnih stroškov',f"=IF({B['hours']}>0,{B['fixed']}/{B['hours']}*F47,0)",EUR),
(51,'Ocena davka',f"=IF({Z},0,F42*{B['tax']}/100)",EUR),
(52,'PREDVIDENI ZASLUŽEK',f'=IF({Z},0,F42-F49-F50-F51)',EUR),
(53,'Zaslužek na uro dela','=IF(F47>0,F52/F47,0)','0.00 "€/h"'),
(54,'Opozorilo',f"=IF(AND(NOT({Z}),F47>0,F53<{B['target']}),\"Pod ciljno postavko!\",\"OK\")",None),
]
k['E11']='Izračun'; k['E11'].font=Font(bold=True,size=12,color=BLUE)
for r,l,f,fm in calc:
    k.cell(r,5,l); c=k.cell(r,6,f)
    if fm: c.number_format=fm
    c.border=Border(bottom=thin)
    if l.isupper() or l.startswith('KONČNA'): k.cell(r,5).font=Font(bold=True); c.font=Font(bold=True)
for lbl,r in [('1. Stroški tiskanja',19),('2. Prodajna vrednost dela',31),('3.–5. Od prihodka do zaslužka',46)]:
    k.cell(r,5,lbl).font=Font(bold=True,color='64748B')
k['F54'].font=Font(bold=True,color='EA580C')

# result boxes
boxes=[('E3','E4','KONČNA CENA','=F44',GREEN),('G3','G4','CENA NA KOS','=F45',BLUE),
('E6','E7','STROŠKI TISKANJA','=F26','0F172A'),('G6','G7','VREDNOST STORITEV','=IF($C$7="Zase",F48,F35+F36)','0F172A'),
('E9','E10','PREDVIDENI ZASLUŽEK','=F52',GREEN),('G9','G10','ZASLUŽEK NA URO','=F53',BLUE)]
box=PatternFill('solid',fgColor='F0FDF4')
for a,b,l,f,col in boxes:
    ra,rb=int(a[1:]),int(b[1:]); c0=a[0]; c1=chr(ord(c0)+1)
    k.merge_cells(f'{c0}{ra}:{c1}{ra}'); k.merge_cells(f'{c0}{rb}:{c1}{rb}')
    k[a]=l; k[a].font=Font(bold=True,size=9,color='64748B'); k[a].fill=box
    k[b]=f; k[b].font=Font(bold=True,size=18,color=col); k[b].fill=box; k[b].number_format=EUR if 'URO' not in l else '0.00 "€/h"'
    k[a].alignment=k[b].alignment=Alignment(horizontal='center')
    k.row_dimensions[rb].height=28
k['E5']='=IF($C$7="Zase","Zase: prikazani so dejanski stroški",IF(F41,"Upoštevano minimalno naročilo",IF('+B['vat']+'="DA","Z DDV","DDV ni obračunan")))'
k['E5'].font=Font(italic=True,size=9,color='64748B')

# chart data
k['J30']='Struktura cene'; k['J30'].font=Font(bold=True)
pie=[('Material','=F20'),('Elektrika','=F21'),('Amortizacija','=F22'),('Potrošni material','=F23'),('Rezerva','=F25'),
('Aktivno delo','=F35+F36'),('Pribitek','=F28'),('Prevoz in ostalo','=F37+F38')]
for i,(l,f) in enumerate(pie):
    k.cell(31+i,10,l); k.cell(31+i,11,f).number_format=EUR
k['J41']='Cena in stroški'; k['J41'].font=Font(bold=True)
bar=[('Dejanski stroški','=F49'),('Fiksni + davek','=F50+F51'),('Zaslužek','=MAX(0,F52)'),('Prihodek','=F42')]
for i,(l,f) in enumerate(bar):
    k.cell(42+i,10,l); k.cell(42+i,11,f).number_format=EUR
k.column_dimensions['J'].width=18; k.column_dimensions['K'].width=12
pc_=PieChart(); pc_.title='Struktura cene'
pc_.add_data(Reference(k,min_col=11,min_row=31,max_row=38)); pc_.set_categories(Reference(k,min_col=10,min_row=31,max_row=38))
pc_.dataLabels=DataLabelList(); pc_.dataLabels.showPercent=True; pc_.height=9; pc_.width=13
k.add_chart(pc_,'J3')
bc=BarChart(); bc.type='bar'; bc.title='Prodajna cena in stroški'; bc.legend=None
bc.add_data(Reference(k,min_col=11,min_row=42,max_row=45)); bc.set_categories(Reference(k,min_col=10,min_row=42,max_row=45))
bc.height=7; bc.width=13; k.add_chart(bc,'J49')
k.freeze_panes='A3'

# ---------- Projekti ----------
pj=wb.create_sheet('Projekti')
header(pj,'Projekti',{c:w for c,w in zip('ABCDEFGHIJKLM',[12,30,14,11,12,13,13,13,10,10,11,11,24])})
thead(pj,3,['Datum','Projekt','Stranka','Način','Status','Končna cena','Stroški','Zaslužek','Ure ocena','Ure dejansko','Material g ocena','Material g dejansko','Opomba'])
pj.append([date(2026,10,25),'Nosilci lučk – 2 leva + 2 desna','Demo','Stranka','Osnutek',132.22,35.49,72.58,3.5,None,768,None,'demo'])
dvS=DataValidation(type='list',formula1='"Osnutek,Ponudba,Potrjeno,V izdelavi,Zaključeno"'); pj.add_data_validation(dvS); dvS.add('E4:E200')
dvM=DataValidation(type='list',formula1='"Stranka,Prijatelj,Zase"'); pj.add_data_validation(dvM); dvM.add('D4:D200')
for r in range(4,60):
    pj.cell(r,1).number_format='DD.MM.YYYY'
    for col in (6,7,8): pj.cell(r,col).number_format=EUR
pj['A62']='Vnesi vrednosti iz Kalkulatorja (Kopiraj → Prilepi kot vrednosti), da se zgodovinske ponudbe ne spremenijo.'; pj['A62'].font=Font(italic=True,color='64748B')
pj.freeze_panes='A4'

for ws in wb.worksheets: ws.sheet_properties.tabColor={'Kalkulator':'F97316'}.get(ws.title,'0F172A')
wb.save('/workspace/hs-3d-kalkulator/excel/HS_3D_Kalkulator_v5.xlsx')
